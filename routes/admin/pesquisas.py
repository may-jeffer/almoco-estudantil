# -*- coding: utf-8 -*-
from flask import render_template, request, redirect, url_for, session, flash, jsonify, Response, send_file
import json
import uuid
import re
from io import BytesIO
from datetime import datetime

from database import closing, get_db_connection, get_config
from utils.auth import is_logged_in_admin, tem_permissao
from utils.helpers import datetime_now_str, date_hoje_str, registrar_auditoria
from utils.qrcode_gen import generate_qr_b64
from . import admin_bp

@admin_bp.route('/admin/pesquisas')
def admin_pesquisas():
    """Listagem de todos os formulários de pesquisa."""
    if not is_logged_in_admin(): return redirect(url_for('admin.admin_login'))
    if not tem_permissao('pesquisas'): return redirect(url_for('admin.admin_dashboard'))

    with closing(get_db_connection()) as conn:
        formularios = conn.execute("""
            SELECT f.*,
                   (SELECT COUNT(*) FROM formulario_perguntas p WHERE p.formulario_id = f.id) as total_perguntas,
                   (SELECT COUNT(*) FROM formulario_respostas_envios e WHERE e.formulario_id = f.id) as total_respostas
            FROM formularios_pesquisa f
            ORDER BY f.id DESC
        """).fetchall()

        total_formularios = len(formularios)
        total_ativos = sum(1 for f in formularios if f['ativo'] == 1)
        total_respostas_geral = sum(f['total_respostas'] for f in formularios)

    hoje_str = date_hoje_str()
    return render_template(
        'admin/pesquisas.html',
        formularios=formularios,
        total_formularios=total_formularios,
        total_ativos=total_ativos,
        total_respostas_geral=total_respostas_geral,
        hoje_str=hoje_str
    )

@admin_bp.route('/admin/pesquisas/novo', methods=['GET', 'POST'])
def admin_pesquisa_novo():
    """Criação de novo formulário de pesquisa (estilo Google Forms)."""
    if not is_logged_in_admin(): return redirect(url_for('admin.admin_login'))
    if not tem_permissao('pesquisas'): return redirect(url_for('admin.admin_dashboard'))

    if request.method == 'POST':
        titulo = request.form.get('titulo', '').strip()
        descricao = request.form.get('descricao', '').strip()
        is_anonimo = 1 if request.form.get('is_anonimo') else 0
        ativo = 1 if request.form.get('ativo') else 0
        data_inicio = request.form.get('data_inicio') or None
        data_fim = request.form.get('data_fim') or None
        perguntas_raw = request.form.get('perguntas_json', '[]')

        if not titulo:
            flash('Informe o título do formulário.', 'error')
            return redirect(url_for('admin.admin_pesquisa_novo'))

        try:
            perguntas = json.loads(perguntas_raw)
        except Exception:
            perguntas = []

        if not perguntas:
            flash('Adicione pelo menos uma pergunta ao formulário.', 'error')
            return redirect(url_for('admin.admin_pesquisa_novo'))

        slug = uuid.uuid4().hex[:10].upper()
        criado_por = session.get('admin_usuario', 'admin')
        agora = datetime_now_str()

        with closing(get_db_connection()) as conn:
            cur = conn.execute("""
                INSERT INTO formularios_pesquisa
                (titulo, descricao, slug, is_anonimo, ativo, data_inicio, data_fim, criado_por, data_criacao)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (titulo, descricao, slug, is_anonimo, ativo, data_inicio, data_fim, criado_por, agora))
            formulario_id = cur.lastrowid

            for idx, p in enumerate(perguntas):
                titulo_p = p.get('titulo', '').strip()
                if not titulo_p: continue
                tipo_p = p.get('tipo', 'multipla_escolha')
                opcoes = p.get('opcoes', [])
                opcoes_json = json.dumps(opcoes, ensure_ascii=False) if tipo_p in ('multipla_escolha', 'checkbox') else None
                obrigatoria = 1 if p.get('obrigatoria') else 0
                ordem = idx

                conn.execute("""
                    INSERT INTO formulario_perguntas
                    (formulario_id, titulo_pergunta, tipo_pergunta, opcoes_json, obrigatoria, ordem)
                    VALUES (?, ?, ?, ?, ?, ?)
                """, (formulario_id, titulo_p, tipo_p, opcoes_json, obrigatoria, ordem))

            conn.commit()
            registrar_auditoria("Criar Pesquisa", f"Criou o formulário '{titulo}' (ID: {formulario_id}, Slug: {slug})")
            flash('Formulário de pesquisa criado com sucesso!', 'success')
            return redirect(url_for('admin.admin_pesquisas'))

    return render_template('admin/pesquisa_form.html', formulario=None, perguntas=[])

@admin_bp.route('/admin/pesquisas/editar/<int:id>', methods=['GET', 'POST'])
def admin_pesquisa_editar(id):
    """Edição de formulário de pesquisa existente."""
    if not is_logged_in_admin(): return redirect(url_for('admin.admin_login'))
    if not tem_permissao('pesquisas'): return redirect(url_for('admin.admin_dashboard'))

    with closing(get_db_connection()) as conn:
        formulario = conn.execute("SELECT * FROM formularios_pesquisa WHERE id = ?", (id,)).fetchone()
        if not formulario:
            flash('Formulário não encontrado.', 'error')
            return redirect(url_for('admin.admin_pesquisas'))

        if request.method == 'POST':
            titulo = request.form.get('titulo', '').strip()
            descricao = request.form.get('descricao', '').strip()
            is_anonimo = 1 if request.form.get('is_anonimo') else 0
            ativo = 1 if request.form.get('ativo') else 0
            data_inicio = request.form.get('data_inicio') or None
            data_fim = request.form.get('data_fim') or None
            perguntas_raw = request.form.get('perguntas_json', '[]')

            if not titulo:
                flash('Informe o título do formulário.', 'error')
                return redirect(url_for('admin.admin_pesquisa_editar', id=id))

            try:
                perguntas = json.loads(perguntas_raw)
            except Exception:
                perguntas = []

            if not perguntas:
                flash('O formulário deve conter pelo menos uma pergunta.', 'error')
                return redirect(url_for('admin.admin_pesquisa_editar', id=id))

            conn.execute("""
                UPDATE formularios_pesquisa
                SET titulo = ?, descricao = ?, is_anonimo = ?, ativo = ?, data_inicio = ?, data_fim = ?
                WHERE id = ?
            """, (titulo, descricao, is_anonimo, ativo, data_inicio, data_fim, id))

            # Atualizar perguntas: remove as antigas e insere a nova estrutura
            conn.execute("DELETE FROM formulario_perguntas WHERE formulario_id = ?", (id,))
            for idx, p in enumerate(perguntas):
                titulo_p = p.get('titulo', '').strip()
                if not titulo_p: continue
                tipo_p = p.get('tipo', 'multipla_escolha')
                opcoes = p.get('opcoes', [])
                opcoes_json = json.dumps(opcoes, ensure_ascii=False) if tipo_p in ('multipla_escolha', 'checkbox') else None
                obrigatoria = 1 if p.get('obrigatoria') else 0
                ordem = idx

                conn.execute("""
                    INSERT INTO formulario_perguntas
                    (formulario_id, titulo_pergunta, tipo_pergunta, opcoes_json, obrigatoria, ordem)
                    VALUES (?, ?, ?, ?, ?, ?)
                """, (id, titulo_p, tipo_p, opcoes_json, obrigatoria, ordem))

            conn.commit()
            registrar_auditoria("Editar Pesquisa", f"Editou o formulário ID {id} ('{titulo}')")
            flash('Formulário atualizado com sucesso!', 'success')
            return redirect(url_for('admin.admin_pesquisas'))

        # GET: carregar perguntas existentes
        perguntas_db = conn.execute("SELECT * FROM formulario_perguntas WHERE formulario_id = ? ORDER BY ordem ASC, id ASC", (id,)).fetchall()
        perguntas = []
        for p in perguntas_db:
            opcoes = json.loads(p['opcoes_json']) if p['opcoes_json'] else []
            perguntas.append({
                'id': p['id'],
                'titulo': p['titulo_pergunta'],
                'tipo': p['tipo_pergunta'],
                'opcoes': opcoes,
                'obrigatoria': bool(p['obrigatoria'])
            })

    return render_template('admin/pesquisa_form.html', formulario=formulario, perguntas=perguntas)

@admin_bp.route('/admin/pesquisas/toggle/<int:id>', methods=['POST'])
def admin_pesquisa_toggle(id):
    """Ativa ou desativa um formulário instantaneamente."""
    if not is_logged_in_admin(): return jsonify({'success': False, 'message': 'Não autorizado'}), 401
    if not tem_permissao('pesquisas'): return jsonify({'success': False, 'message': 'Sem permissão'}), 403

    with closing(get_db_connection()) as conn:
        form = conn.execute("SELECT id, titulo, ativo FROM formularios_pesquisa WHERE id = ?", (id,)).fetchone()
        if not form:
            return jsonify({'success': False, 'message': 'Formulário não encontrado'}), 404

        novo_status = 0 if form['ativo'] == 1 else 1
        conn.execute("UPDATE formularios_pesquisa SET ativo = ? WHERE id = ?", (novo_status, id))
        conn.commit()

        status_txt = "ativado" if novo_status == 1 else "desativado"
        registrar_auditoria("Toggle Pesquisa", f"O formulário '{form['titulo']}' foi {status_txt}")

    return jsonify({'success': True, 'novo_status': novo_status, 'message': f"Formulário {status_txt} com sucesso."})

@admin_bp.route('/admin/pesquisas/excluir/<int:id>', methods=['POST'])
def admin_pesquisa_excluir(id):
    """Exclui um formulário e todas as suas respostas associadas."""
    if not is_logged_in_admin(): return redirect(url_for('admin.admin_login'))
    if not tem_permissao('pesquisas'): return redirect(url_for('admin.admin_dashboard'))

    with closing(get_db_connection()) as conn:
        form = conn.execute("SELECT id, titulo FROM formularios_pesquisa WHERE id = ?", (id,)).fetchone()
        if form:
            # Excluir itens de respostas, envios, perguntas e formulário
            envios = conn.execute("SELECT id FROM formulario_respostas_envios WHERE formulario_id = ?", (id,)).fetchall()
            envio_ids = [e['id'] for e in envios]
            if envio_ids:
                ph = ','.join(['?'] * len(envio_ids))
                conn.execute(f"DELETE FROM formulario_respostas_itens WHERE envio_id IN ({ph})", envio_ids)
            conn.execute("DELETE FROM formulario_respostas_envios WHERE formulario_id = ?", (id,))
            conn.execute("DELETE FROM formulario_perguntas WHERE formulario_id = ?", (id,))
            conn.execute("DELETE FROM formularios_pesquisa WHERE id = ?", (id,))
            conn.commit()

            registrar_auditoria("Excluir Pesquisa", f"Excluiu o formulário '{form['titulo']}' (ID {id})")
            flash(f"Formulário '{form['titulo']}' excluído com sucesso.", 'success')
        else:
            flash('Formulário não encontrado.', 'error')

    return redirect(url_for('admin.admin_pesquisas'))

@admin_bp.route('/admin/pesquisas/respostas/<int:id>')
def admin_pesquisa_respostas(id):
    """Dashboard analítico das respostas do formulário."""
    if not is_logged_in_admin(): return redirect(url_for('admin.admin_login'))
    if not tem_permissao('pesquisas'): return redirect(url_for('admin.admin_dashboard'))

    with closing(get_db_connection()) as conn:
        formulario = conn.execute("SELECT * FROM formularios_pesquisa WHERE id = ?", (id,)).fetchone()
        if not formulario:
            flash('Formulário não encontrado.', 'error')
            return redirect(url_for('admin.admin_pesquisas'))

        perguntas = conn.execute("SELECT * FROM formulario_perguntas WHERE formulario_id = ? ORDER BY ordem ASC, id ASC", (id,)).fetchall()
        envios = conn.execute("""
            SELECT e.*, a.nome as aluno_nome_db, a.matricula as aluno_matricula_db, t.nome as turma_nome
            FROM formulario_respostas_envios e
            LEFT JOIN alunos a ON e.aluno_id = a.id
            LEFT JOIN turmas t ON a.turma_id = t.id
            WHERE e.formulario_id = ?
            ORDER BY e.id DESC
        """, (id,)).fetchall()

        itens = conn.execute("""
            SELECT i.*, p.tipo_pergunta
            FROM formulario_respostas_itens i
            JOIN formulario_perguntas p ON i.pergunta_id = p.id
            WHERE p.formulario_id = ?
        """, (id,)).fetchall()

    # Mapear respostas por envio_id e pergunta_id
    respostas_map = {}
    for it in itens:
        envio_id = it['envio_id']
        pergunta_id = it['pergunta_id']
        if envio_id not in respostas_map:
            respostas_map[envio_id] = {}
        respostas_map[envio_id][pergunta_id] = it['resposta_texto']

    # Compilar estatísticas por pergunta
    estatisticas = []
    for p in perguntas:
        p_id = p['id']
        tipo = p['tipo_pergunta']
        opcoes = json.loads(p['opcoes_json']) if p['opcoes_json'] else []
        respostas_pergunta = []

        for e in envios:
            ans = respostas_map.get(e['id'], {}).get(p_id)
            if ans is not None and ans.strip():
                respostas_pergunta.append(ans)

        stat = {
            'id': p_id,
            'titulo': p['titulo_pergunta'],
            'tipo': tipo,
            'total_respostas': len(respostas_pergunta),
            'opcoes': opcoes,
            'contagem_opcoes': {},
            'respostas_texto': []
        }

        if tipo in ('multipla_escolha', 'checkbox'):
            contagem = {op: 0 for op in opcoes}
            for ans in respostas_pergunta:
                if tipo == 'checkbox':
                    try:
                        selecionados = json.loads(ans) if ans.startswith('[') else [x.strip() for x in ans.split(',')]
                    except Exception:
                        selecionados = [x.strip() for x in ans.split(',')]
                    for sel in selecionados:
                        if sel in contagem:
                            contagem[sel] += 1
                        else:
                            contagem[sel] = 1
                else:
                    if ans in contagem:
                        contagem[ans] += 1
                    else:
                        contagem[ans] = 1
            stat['contagem_opcoes'] = contagem
        else:
            stat['respostas_texto'] = respostas_pergunta

        estatisticas.append(stat)

    return render_template(
        'admin/pesquisa_respostas.html',
        formulario=formulario,
        perguntas=perguntas,
        envios=envios,
        respostas_map=respostas_map,
        estatisticas=estatisticas,
        total_envios=len(envios)
    )

@admin_bp.route('/admin/pesquisas/excel/<int:id>')
def admin_pesquisa_excel(id):
    """Exporta todas as respostas da pesquisa para uma planilha Excel (.xlsx)."""
    if not is_logged_in_admin(): return redirect(url_for('admin.admin_login'))
    if not tem_permissao('pesquisas'): return redirect(url_for('admin.admin_dashboard'))

    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    with closing(get_db_connection()) as conn:
        formulario = conn.execute("SELECT * FROM formularios_pesquisa WHERE id = ?", (id,)).fetchone()
        if not formulario:
            flash('Formulário não encontrado.', 'error')
            return redirect(url_for('admin.admin_pesquisas'))

        perguntas = conn.execute("SELECT * FROM formulario_perguntas WHERE formulario_id = ? ORDER BY ordem ASC, id ASC", (id,)).fetchall()
        envios = conn.execute("""
            SELECT e.*, a.nome as aluno_nome_db, a.matricula as aluno_matricula_db, t.nome as turma_nome
            FROM formulario_respostas_envios e
            LEFT JOIN alunos a ON e.aluno_id = a.id
            LEFT JOIN turmas t ON a.turma_id = t.id
            WHERE e.formulario_id = ?
            ORDER BY e.id ASC
        """, (id,)).fetchall()

        itens = conn.execute("""
            SELECT i.*
            FROM formulario_respostas_itens i
            JOIN formulario_perguntas p ON i.pergunta_id = p.id
            WHERE p.formulario_id = ?
        """, (id,)).fetchall()

    respostas_map = {}
    for it in itens:
        envio_id = it['envio_id']
        p_id = it['pergunta_id']
        if envio_id not in respostas_map:
            respostas_map[envio_id] = {}
        respostas_map[envio_id][p_id] = it['resposta_texto']

    wb = Workbook()
    ws = wb.active
    ws.title = "Respostas Pesquisa"
    ws.views.sheetView[0].showGridLines = True

    thin_side = Side(style='thin', color='CBD5E1')
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    header_font = Font(name='Segoe UI', bold=True, color='FFFFFF', size=11)

    config = get_config()
    ws['A1'] = f"RELATÓRIO DE PESQUISA: {formulario['titulo'].upper()}"
    ws['A1'].font = Font(name='Segoe UI', bold=True, size=14, color='1E293B')
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')

    tipo_anonimo_txt = "Anônima (Sem identificação)" if formulario['is_anonimo'] == 1 else "Identificada (Alunos Registrados)"
    ws['A2'] = f"Tipo: {tipo_anonimo_txt} | Total de Respostas: {len(envios)} | Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws['A2'].font = Font(name='Segoe UI', italic=True, size=9, color='64748B')

    headers = ["#", "Data / Hora"]
    if formulario['is_anonimo'] == 0:
        headers.extend(["Estudante", "Matrícula", "Turma"])

    for p in perguntas:
        headers.append(p['titulo_pergunta'])

    ws.append([]) # Linha 3 vazia
    ws.append(headers) # Linha 4 Cabeçalho
    ws.row_dimensions[4].height = 28

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    # Dados
    for idx, e in enumerate(envios, 1):
        row_data = [idx, e['data_envio']]
        if formulario['is_anonimo'] == 0:
            nome = e['aluno_nome_db'] or e['aluno_identificacao'] or '—'
            mat = e['aluno_matricula_db'] or '—'
            turma = e['turma_nome'] or '—'
            row_data.extend([nome, mat, turma])

        for p in perguntas:
            val = respostas_map.get(e['id'], {}).get(p['id'], '')
            row_data.append(val or '')

        ws.append(row_data)
        row_num = ws.max_row
        ws.row_dimensions[row_num].height = 20

        for col_num in range(1, len(headers) + 1):
            c = ws.cell(row=row_num, column=col_num)
            c.border = thin_border
            c.alignment = Alignment(horizontal='left' if col_num > 2 else 'center', vertical='center')

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or '')
            if cell.row > 2:
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    slug_safe = re.sub(r'[^a-zA-Z0-9_-]', '_', formulario['titulo'][:30])
    filename = f"pesquisa_{slug_safe}_{date_hoje_str()}.xlsx"

    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@admin_bp.route('/admin/pesquisas/qrcode/<int:id>')
def admin_pesquisa_qrcode(id):
    """Gera o QR Code da pesquisa em base64 para exibição e download."""
    if not is_logged_in_admin(): return jsonify({'success': False}), 401
    if not tem_permissao('pesquisas'): return jsonify({'success': False}), 403

    with closing(get_db_connection()) as conn:
        formulario = conn.execute("SELECT id, titulo, slug FROM formularios_pesquisa WHERE id = ?", (id,)).fetchone()
        if not formulario:
            return jsonify({'success': False, 'message': 'Formulário não encontrado'}), 404

    url_publica = request.host_url.rstrip('/') + url_for('main.pesquisa_responder', slug=formulario['slug'])
    qr_b64 = generate_qr_b64(url_publica)

    return jsonify({
        'success': True,
        'titulo': formulario['titulo'],
        'slug': formulario['slug'],
        'url': url_publica,
        'qr_b64': qr_b64
    })
